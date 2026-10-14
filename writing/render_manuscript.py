"""Render the adjacent Markdown manuscript as a Word reading copy.

Run with the Codex bundled Python runtime. Markdown is the canonical source.
Style: narrative_proposal; named overrides: A4 page, Korean font,
editorial cover title (24pt), chapter/outline page starts, compact intro/outline.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Mm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / '책_집필_원고.md'
OUTPUT = ROOT / '책_집필_원고.docx'
FONT = 'Apple SD Gothic Neo'

def style(doc, name, size, color, before, after, bold=False):
    s = doc.styles[name]
    s.font.name = FONT
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = bold
    fonts = s.element.get_or_add_rPr().get_or_add_rFonts()
    for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + k), FONT)
    pf = s.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 320 / 240
    pf.widow_control = True
    pf.keep_with_next = name.startswith('Heading')
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    s.element.get_or_add_pPr().append(snap)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'ko-KR')
    lang.set(qn('w:eastAsia'), 'ko-KR')
    s.element.get_or_add_rPr().append(lang)
    return s

def build():
    text = SOURCE.read_text()
    assert len(re.findall(r'^## [1-8]장\.', text, re.M)) == 8
    assert '```' not in text
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(.492)
    grid = sec._sectPr.find(qn('w:docGrid'))
    if grid is not None:
        grid.set(qn('w:type'), 'default')
        grid.attrib.pop(qn('w:linePitch'), None)

    normal = style(doc, 'Normal', 11, '202020', 0, 8)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style(doc, 'Title', 24, '0B2545', 12, 14, True)
    style(doc, 'Subtitle', 11, '555555', 0, 16)
    style(doc, 'Heading 1', 16, '2E74B5', 18, 10, True)
    style(doc, 'Heading 2', 13, '2E74B5', 12, 6, True)
    style(doc, 'Heading 3', 12, '1F4D78', 8, 4, True)
    doc.styles.add_style('Intro Body', WD_STYLE_TYPE.PARAGRAPH)
    doc.styles.add_style('Outline Body', WD_STYLE_TYPE.PARAGRAPH)
    intro_body = style(doc, 'Intro Body', 10.5, '202020', 0, 6)
    intro_body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_body.paragraph_format.line_spacing = 1.25
    outline_body = style(doc, 'Outline Body', 10.5, '202020', 0, 5)
    outline_body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    outline_body.paragraph_format.line_spacing = 1.15
    for n in ('Header', 'Footer'):
        style(doc, n, 9, '666666', 0, 0)
    # python-docx's built-in Title style may inherit a template border.
    # The editorial-cover override uses typography only.
    title_ppr = doc.styles['Title'].element.get_or_add_pPr()
    border = title_ppr.find(qn('w:pBdr'))
    if border is not None:
        title_ppr.remove(border)

    header = sec.header.paragraphs[0]
    header.text = '풀스택 서비스 개발의 기초 | 집필 초안'
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('집필 원고 v0.1  ·  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    in_intro = False
    in_outline = False
    intro_chars = []
    chars = 0
    for block in re.split(r'\n\s*\n', text.strip()):
        if block.startswith('# '):
            p = doc.add_paragraph(block[2:], 'Title')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.startswith('## '):
            title = block[3:]
            p = doc.add_paragraph(title, 'Heading 1')
            in_outline = False
            if re.match(r'[1-8]장\.', title):
                p.paragraph_format.page_break_before = True
        elif block.startswith('### '):
            title = block[4:]
            if in_intro:
                intro_chars.append(chars)
                chars = 0
            in_intro = title == '장 도입문'
            in_outline = title == '세부 아웃라인'
            p = doc.add_paragraph(title, 'Heading 2')
            if title == '세부 아웃라인':
                p.paragraph_format.page_break_before = True
        elif block.startswith('#### '):
            doc.add_paragraph(block[5:], 'Heading 3')
        else:
            if in_intro:
                chars += len(block)
            if block.startswith('장별 아웃라인과 도입문 |'):
                name = 'Subtitle'
            elif in_intro:
                name = 'Intro Body'
            elif in_outline:
                name = 'Outline Body'
            else:
                name = 'Normal'
            doc.add_paragraph(block.replace('\n', ' '), name)

    doc.core_properties.title = '예제로 따라해보는 풀스택 서비스 개발의 기초'
    doc.core_properties.subject = '장별 아웃라인과 도입문'
    doc.core_properties.author = ''
    doc.core_properties.keywords = '집필 초안, 풀스택, 입문, v0.1'
    doc.save(OUTPUT)
    assert len(intro_chars) == 8
    print('Intro characters by chapter:', intro_chars)
    print('Written:', OUTPUT)

if __name__ == '__main__':
    build()
